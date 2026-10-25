import docx

doc = docx.Document('Final_BE_REPORT_2026 (1).docx')
for s_idx, section in enumerate(doc.sections):
    print(f"=== Section {s_idx} ===")
    header = section.header
    footer = section.footer
    print("  Header text:", [p.text for p in header.paragraphs if p.text.strip()])
    print("  Footer text:", [p.text for p in footer.paragraphs if p.text.strip()])
