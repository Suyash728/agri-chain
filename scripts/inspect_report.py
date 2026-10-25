import docx

doc = docx.Document('Final_BE_REPORT_2026 (1).docx')
p_sample = [0, 1, 16, 28, 54, 103, 104, 105, 115, 150, 155]
for idx in p_sample:
    if idx < len(doc.paragraphs):
        p = doc.paragraphs[idx]
        pf = p.paragraph_format
        r_font = p.runs[0].font.name if p.runs and p.runs[0].font.name else 'None'
        r_size = p.runs[0].font.size.pt if p.runs and p.runs[0].font.size else 'None'
        print(f"P{idx} [{p.style.name}]: text='{p.text[:35]}' | align={p.alignment} | left_indent={pf.left_indent} | right_indent={pf.right_indent} | space_before={pf.space_before} | space_after={pf.space_after} | font={r_font} {r_size}")
