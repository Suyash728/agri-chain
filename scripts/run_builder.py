# scripts/run_builder.py
import os
import sys

# Ensure root directory is in sys.path
ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)

import docx
from scripts.report_helpers import init_doc
from scripts.section_front_matter import build_front_matter
from scripts.section_chapters_1_3 import build_chapters_1_3
from scripts.section_chapter_4 import build_chapter_4
from scripts.section_chapter_5_6 import build_chapters_5_6

def build_full_report():
    print("=" * 60)
    print("STARTING FULL BE MAJOR PROJECT REPORT BUILD")
    print("=" * 60)

    # 1. Initialize Document
    print("[1/5] Initializing Document with standard 1.0\" margins and college headers...")
    doc = init_doc()

    # 2. Front Matter
    print("[2/5] Assembling Front Matter (Cover, Approval, Cert, Decl, Ack, Abstract, TOC, LOF, LOT, Abbreviations)...")
    build_front_matter(doc)

    # 3. Chapters 1 to 3
    print("[3/5] Assembling Chapter 1 (Intro), Chapter 2 (Literature Review + Tables 2.1-2.5), Chapter 3 (Requirements)...")
    build_chapters_1_3(doc)

    # 4. Chapter 4
    print("[4/5] Assembling Chapter 4 (Methodology, Figs 4.1-4.7, Tables 4.1-4.2)...")
    build_chapter_4(doc)

    # 5. Chapters 5 & 6, Appendices, Bibliography
    print("[5/5] Assembling Chapter 5 (Results, Figs 5.1-5.11, Tables 5.1-5.4), Chapter 6 (Conclusion), Appendices, Bibliography...")
    build_chapters_5_6(doc)

    # Output targets
    target_files = [
        os.path.join(ROOT_DIR, "Final_BE_Report_AgriChain.docx"),
        os.path.join(ROOT_DIR, "Final_BE_REPORT_2026 (1).docx")
    ]

    for target in target_files:
        doc.save(target)
        print(f"[OK] Saved generated report to: {target}")

    # Post-generation verification
    print("\n" + "=" * 60)
    print("RUNNING POST-BUILD VERIFICATION")
    print("=" * 60)

    verify_doc = docx.Document(target_files[0])
    num_paras = len(verify_doc.paragraphs)
    num_tables = len(verify_doc.tables)
    num_sections = len(verify_doc.sections)

    sec0 = verify_doc.sections[0]
    margins = (
        f"top={sec0.top_margin.inches:.2f}\", "
        f"bottom={sec0.bottom_margin.inches:.2f}\", "
        f"left={sec0.left_margin.inches:.2f}\", "
        f"right={sec0.right_margin.inches:.2f}\""
    )

    missing = [p.text for p in verify_doc.paragraphs if "not found" in p.text.lower()]
    fig_caps = [p.text for p in verify_doc.paragraphs if p.text.startswith("Figure ") and ":" in p.text]
    table_caps = [p.text for p in verify_doc.paragraphs if p.text.startswith("Table ") and ":" in p.text]

    print(f"Total Sections  : {num_sections}")
    print(f"Page Margins    : {margins}")
    print(f"Total Paragraphs: {num_paras}")
    print(f"Total Tables    : {num_tables}")
    print(f"Figure Captions : {len(fig_caps)}")
    print(f"Table Captions  : {len(table_caps)}")
    print(f"Missing Images  : {len(missing)}")

    if missing:
        print("[WARNING] Missing image references detected:", missing)
    else:
        print("[SUCCESS] All diagrams and figures resolved and successfully embedded!")

    print("=" * 60)
    print("BUILD AND VERIFICATION COMPLETED SUCCESSFULLY!")
    print("=" * 60)

if __name__ == "__main__":
    build_full_report()
