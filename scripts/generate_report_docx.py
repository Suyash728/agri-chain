# scripts/generate_report_docx.py
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

import scripts.report_data as rd

def create_report():
    doc = docx.Document()
    
    # Page setup - Standard A4 with normal 1.0 inch margins
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.header_distance = Inches(0.4)
        section.footer_distance = Inches(0.4)
        
        # Header setup
        header = section.header
        header.is_linked_to_previous = False
        hp1 = header.paragraphs[0]
        hp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp1.paragraph_format.space_after = Pt(2)
        r1 = hp1.add_run("Reg. No. E-91,(Osmanabad) Dated 30-09-80                                              DTE INST. CODE : 3190")
        r1.font.name = "Georgia"
        r1.font.size = Pt(8.5)
        r1.font.color.rgb = RGBColor(100, 100, 100)
        
        hp2 = header.add_paragraph()
        hp2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp2.paragraph_format.space_after = Pt(4)
        r2 = hp2.add_run("Department of Artificial Intelligence (AI) and Data Science")
        r2.font.name = "Georgia"
        r2.font.size = Pt(8.5)
        r2.font.bold = True
        r2.font.color.rgb = RGBColor(70, 70, 70)

    # Styling helper functions
    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def set_table_borders(table, color="D0D0D0", sz="4", val="single"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.left_indent = Inches(0)
        p.paragraph_format.right_indent = Inches(0)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Georgia"
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = RGBColor(20, 45, 85)
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0)
        p.paragraph_format.right_indent = Inches(0)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Georgia"
        r.font.size = Pt(13.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(30, 50, 70)
        return p

    def add_heading_3(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Inches(0)
        p.paragraph_format.right_indent = Inches(0)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Georgia"
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = RGBColor(40, 60, 80)
        return p

    def add_para(text, bold_prefix="", italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=5):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.left_indent = Inches(0)
        p.paragraph_format.right_indent = Inches(0)
        
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = "Georgia"
            rb.font.size = Pt(12)
            rb.font.bold = True
            rb.font.color.rgb = RGBColor(30, 30, 30)
            
        r = p.add_run(text)
        r.font.name = "Georgia"
        r.font.size = Pt(12)
        r.font.italic = italic
        r.font.color.rgb = RGBColor(40, 40, 40)
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.right_indent = Inches(0)
        
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = "Georgia"
            rb.font.size = Pt(11.5)
            rb.font.bold = True
            
        r = p.add_run(text)
        r.font.name = "Georgia"
        r.font.size = Pt(11.5)
        r.font.color.rgb = RGBColor(40, 40, 40)
        return p

    def add_figure(img_path, caption, width_in=5.8, desc_text=""):
        if desc_text:
            add_para(desc_text)
            
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(6)
            p_img.paragraph_format.space_after = Pt(2)
            p_img.paragraph_format.left_indent = Inches(0)
            p_img.paragraph_format.right_indent = Inches(0)
            p_img.paragraph_format.keep_with_next = True
            run = p_img.add_run()
            run.add_picture(img_path, width=Inches(width_in))
        else:
            p_miss = doc.add_paragraph(f"[Image file not found: {img_path}]")
            p_miss.alignment = WD_ALIGN_PARAGRAPH.CENTER

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(8)
        p_cap.paragraph_format.left_indent = Inches(0)
        p_cap.paragraph_format.right_indent = Inches(0)
        r_cap = p_cap.add_run(caption)
        r_cap.font.name = "Georgia"
        r_cap.font.size = Pt(10.5)
        r_cap.font.bold = True
        r_cap.font.color.rgb = RGBColor(25, 45, 75)

    def add_styled_table(header, rows, col_widths=None, caption=""):
        if caption:
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(8)
            p_cap.paragraph_format.space_after = Pt(3)
            p_cap.paragraph_format.left_indent = Inches(0)
            p_cap.paragraph_format.right_indent = Inches(0)
            p_cap.paragraph_format.keep_with_next = True
            r_cap = p_cap.add_run(caption)
            r_cap.font.name = "Georgia"
            r_cap.font.size = Pt(11)
            r_cap.font.bold = True
            r_cap.font.color.rgb = RGBColor(20, 40, 70)

        table = doc.add_table(rows=len(rows) + 1, cols=len(header))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, color="B8C0C8", sz="4", val="single")

        # Header Row
        hdr_cells = table.rows[0].cells
        for col_idx, heading_text in enumerate(header):
            cell = hdr_cells[col_idx]
            cell.text = heading_text
            set_cell_background(cell, "1E3A2F")  # Deep Forest Green / Academic slate
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for r in p.runs:
                    r.font.name = "Georgia"
                    r.font.size = Pt(9.5)
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(255, 255, 255)

        # Body Rows
        for r_idx, row_data in enumerate(rows):
            row_cells = table.rows[r_idx + 1].cells
            bg_color = "F8FAF9" if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, cell_value in enumerate(row_data):
                cell = row_cells[c_idx]
                cell.text = str(cell_value)
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=80, bottom=80, left=110, right=110)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx > 0 and len(str(cell_value)) > 15 else WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.1
                    for r in p.runs:
                        r.font.name = "Georgia"
                        r.font.size = Pt(9.5)
                        r.font.color.rgb = RGBColor(40, 40, 40)

        # Apply Column Widths if provided
        if col_widths and len(col_widths) == len(header):
            for row in table.rows:
                for idx, w in enumerate(col_widths):
                    row.cells[idx].width = Inches(w)

        # Space after table
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(6)

    print("Framework initialized. Building sections...")
    return doc, add_heading_1, add_heading_2, add_heading_3, add_para, add_bullet, add_figure, add_styled_table

if __name__ == "__main__":
    doc, h1, h2, h3, para, bullet, fig, tbl = create_report()
    doc.save("test_init.docx")
    print("test_init.docx created successfully")
