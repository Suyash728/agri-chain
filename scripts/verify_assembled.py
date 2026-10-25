import docx

doc = docx.Document('test_assembled_report.docx')
print('Total sections:', len(doc.sections))
print('Total paragraphs:', len(doc.paragraphs))
print('Total tables:', len(doc.tables))

for i, sec in enumerate(doc.sections):
    print(f'Section {i}: top={sec.top_margin.inches:.2f}", bottom={sec.bottom_margin.inches:.2f}", left={sec.left_margin.inches:.2f}", right={sec.right_margin.inches:.2f}"')

# Check for missing images
missing_imgs = [p.text for p in doc.paragraphs if 'not found' in p.text.lower()]
print('Missing images:', len(missing_imgs), missing_imgs)

# Check all figure captions
fig_caps = [p.text for p in doc.paragraphs if p.text.startswith('Figure ')]
print(f'Figure captions found ({len(fig_caps)}):')
for fc in fig_caps:
    print('  ', fc)

# Check all table captions
tbl_caps = [p.text for p in doc.paragraphs if p.text.startswith('Table ')]
print(f'Table captions found ({len(tbl_caps)}):')
for tc in tbl_caps:
    print('  ', tc)

# Check paragraph alignments and indents
indents = set()
for p in doc.paragraphs:
    pf = p.paragraph_format
    if pf.left_indent is not None:
        indents.add(f"left: {pf.left_indent.inches:.2f}\"")
    if pf.right_indent is not None:
        indents.add(f"right: {pf.right_indent.inches:.2f}\"")

print('Unique indents detected:', indents)
