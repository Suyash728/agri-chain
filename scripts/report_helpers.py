# scripts/report_helpers.py
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def init_doc():
    doc = docx.Document()
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.header_distance = Inches(0.45)
        section.footer_distance = Inches(0.45)
        section.different_first_page_header_footer = True
        
        # Setup headers (Both normal and first page get the standard 2-line college header)
        for h in [section.header, section.first_page_header]:
            h.is_linked_to_previous = False
            hp1 = h.paragraphs[0]
            hp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            hp1.paragraph_format.space_after = Pt(2)
            hp1.paragraph_format.left_indent = Inches(0)
            hp1.paragraph_format.right_indent = Inches(0)
            r1 = hp1.add_run("Reg. No. E-91,(Osmanabad) Dated 30-09-80                                              DTE INST. CODE : 3190")
            r1.font.name = "Verdana"
            r1.font.size = Pt(8.5)
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(90, 90, 90)
            
            hp2 = h.add_paragraph()
            hp2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            hp2.paragraph_format.space_after = Pt(4)
            hp2.paragraph_format.left_indent = Inches(0)
            hp2.paragraph_format.right_indent = Inches(0)
            r2 = hp2.add_run("Department of Artificial Intelligence (AI) and Data Science")
            r2.font.name = "Georgia"
            r2.font.size = Pt(9.0)
            r2.font.bold = True
            r2.font.color.rgb = RGBColor(60, 60, 60)
        
        # First Page Footer (College Address)
        fp_ftr = section.first_page_footer
        fp_ftr.is_linked_to_previous = False
        fp_p1 = fp_ftr.paragraphs[0]
        fp_p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp_p1.paragraph_format.space_after = Pt(1)
        r_fp1 = fp_p1.add_run("Plot No.12, Sector-22, Phase-II, Nerul, Navi Mumbai - 400 706")
        r_fp1.font.name = "Georgia"
        r_fp1.font.size = Pt(8.5)
        r_fp1.font.color.rgb = RGBColor(100, 100, 100)
        
        fp_p2 = fp_ftr.add_paragraph()
        fp_p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp_p2.paragraph_format.space_after = Pt(1)
        r_fp2 = fp_p2.add_run("Tel. No. 022 61115444 (30 Lines), Fax No. 022 6111 5400  Website : www.terna.org")
        r_fp2.font.name = "Georgia"
        r_fp2.font.size = Pt(8.5)
        r_fp2.font.color.rgb = RGBColor(100, 100, 100)
        
        fp_p3 = fp_ftr.add_paragraph()
        fp_p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp_p3.paragraph_format.space_after = Pt(0)
        r_fp3 = fp_p3.add_run("Email : principal@terna.org")
        r_fp3.font.name = "Georgia"
        r_fp3.font.size = Pt(8.5)
        r_fp3.font.color.rgb = RGBColor(100, 100, 100)
        
        # Subsequent Pages Footer (Centered dynamic page number)
        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.space_before = Pt(4)
        fp.paragraph_format.space_after = Pt(0)
        fp.paragraph_format.left_indent = Inches(0)
        fp.paragraph_format.right_indent = Inches(0)
        
        frun = fp.add_run()
        frun.font.name = "Georgia"
        frun.font.size = Pt(10.0)
        frun.font.color.rgb = RGBColor(80, 80, 80)
        
        fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
        instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
        fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
        fldChar3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
        frun._r.append(fldChar1)
        frun._r.append(instrText)
        frun._r.append(fldChar2)
        frun._r.append(fldChar3)

    return doc

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=130, right=130):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="B8C0C8", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Inches(0)
    p.paragraph_format.right_indent = Inches(0)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Georgia"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(16, 42, 77)
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0)
    p.paragraph_format.right_indent = Inches(0)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Georgia"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(25, 45, 75)
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0)
    p.paragraph_format.right_indent = Inches(0)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Georgia"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(40, 60, 80)
    return p

def add_para(doc, text, bold_prefix="", italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=5):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Inches(0)
    p.paragraph_format.right_indent = Inches(0)
    
    if bold_prefix:
        rb = p.add_run(bold_prefix)
        rb.font.name = "Georgia"
        rb.font.size = Pt(12)
        rb.font.bold = True
        rb.font.color.rgb = RGBColor(20, 20, 20)
        
    r = p.add_run(text)
    r.font.name = "Georgia"
    r.font.size = Pt(12)
    r.font.italic = italic
    r.font.color.rgb = RGBColor(40, 40, 40)
    return p

def add_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0)
    
    if bold_prefix:
        rb = p.add_run(bold_prefix)
        rb.font.name = "Georgia"
        rb.font.size = Pt(11.5)
        rb.font.bold = True
        rb.font.color.rgb = RGBColor(20, 20, 20)
        
    r = p.add_run(text)
    r.font.name = "Georgia"
    r.font.size = Pt(11.5)
    r.font.color.rgb = RGBColor(40, 40, 40)
    return p

def add_figure(doc, img_path, caption, width_in=5.8, desc_text=""):
    if desc_text:
        add_para(doc, desc_text)
        
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(3)
        p_img.paragraph_format.left_indent = Inches(0)
        p_img.paragraph_format.right_indent = Inches(0)
        p_img.paragraph_format.keep_with_next = True
        run = p_img.add_run()
        run.add_picture(img_path, width=Inches(width_in))
    else:
        p_miss = doc.add_paragraph(f"[Image file not found: {img_path}]")
        p_miss.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(3)
    p_cap.paragraph_format.space_after = Pt(10)
    p_cap.paragraph_format.left_indent = Inches(0)
    p_cap.paragraph_format.right_indent = Inches(0)
    r_cap = p_cap.add_run(caption)
    r_cap.font.name = "Georgia"
    r_cap.font.size = Pt(10.5)
    r_cap.font.bold = True
    r_cap.font.color.rgb = RGBColor(25, 45, 75)

def add_styled_table(doc, header, rows, col_widths=None, caption=""):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(10)
        p_cap.paragraph_format.space_after = Pt(4)
        p_cap.paragraph_format.left_indent = Inches(0)
        p_cap.paragraph_format.right_indent = Inches(0)
        p_cap.paragraph_format.keep_with_next = True
        r_cap = p_cap.add_run(caption)
        r_cap.font.name = "Georgia"
        r_cap.font.size = Pt(11)
        r_cap.font.bold = True
        r_cap.font.color.rgb = RGBColor(20, 40, 70)

    table = doc.add_table(rows=len(rows) + 1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="B8C0C8", sz="4", val="single")

    # Header Row
    hdr_cells = table.rows[0].cells
    for col_idx, heading_text in enumerate(header):
        cell = hdr_cells[col_idx]
        cell.text = heading_text
        set_cell_background(cell, "1E3A2F")  # Deep Academic Forest Green
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            for r in p.runs:
                r.font.name = "Georgia"
                r.font.size = Pt(9.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)

    # Body Rows
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAF9" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell_value in enumerate(row_data):
            cell = row_cells[c_idx]
            cell.text = str(cell_value)
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80, left=110, right=110)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx > 0 and len(str(cell_value)) > 15 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.1
                for r in p.runs:
                    r.font.name = "Georgia"
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(40, 40, 40)

    if col_widths and len(col_widths) == len(header):
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)
