import docx

doc = docx.Document('Final_BE_REPORT_2026 (1).docx')
for i, p in enumerate(doc.paragraphs):
    for r in p.runs:
        if 'graphic' in r._r.xml:
            # find blip r:embed
            import xml.etree.ElementTree as ET
            tree = ET.fromstring(r._r.xml)
            for elem in tree.iter():
                if elem.tag.endswith('blip'):
                    embed = elem.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                    target = doc.part.rels[embed].target_ref
                    caption = doc.paragraphs[i+1].text.strip() if i+1 < len(doc.paragraphs) else ''
                    print(f"P{i} embeds {target} -> Caption: {caption[:60]}")
